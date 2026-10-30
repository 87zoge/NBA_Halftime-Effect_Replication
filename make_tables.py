"""
Build output/tables.docx from the CSVs written by analysis.py: the paper's Tables 1-6 and A1,
in Times New Roman with three horizontal rules, standard errors in parentheses, significance
stars and a note under each table.

Called at the end of analysis.py; can also be run on its own once the CSVs exist.
"""

from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

OUT_DIR = Path(__file__).parent / "output"
FONT = "Times New Roman"
STARS = "* p < 0.10, ** p < 0.05, *** p < 0.01."
DASH = "–"
TEXT_WIDTH_CM = 16  # A4 with 2.5 cm margins
Z_TEST = "z-test for two independent estimates, z = (b₁ − b₂) / √(SE₁² + SE₂²)"
SHORT_LINE = (1995, 1997)  # as in analysis.py, printed in bold in Table A1


def num(x: float, d: int = 1) -> str:
    """Fixed decimals, thousands separator, true minus sign."""
    return f"{x:,.{d}f}".replace("-", "−")


def stars(p: float) -> str:
    return "***" if p < 0.01 else "**" if p < 0.05 else "*" if p < 0.1 else ""


def est(r, d: int = 1) -> tuple[str, str]:
    """Estimate with stars, and its SE in parentheses (two table rows)."""
    return num(r["coef"], d) + stars(r["p"]), f"({num(r['se'], d)})"


def era_header(label: str) -> str:
    """'B&P sample 1993/94-2008/09' -> two-line column header."""
    name, seasons = label.rsplit(" ", 1)
    return f"{name}\n{seasons.replace('-', DASH)}"


def read(name: str, index=0) -> pd.DataFrame:
    return pd.read_csv(OUT_DIR / f"{name}.csv", index_col=index)


# Word formatting

def add_rule(cells, side: str, size: int):
    """Horizontal border on one side of each cell (size in eighths of a point)."""
    for cell in cells:
        tc_pr = cell._tc.get_or_add_tcPr()
        borders = tc_pr.find(qn("w:tcBorders"))
        if borders is None:
            borders = OxmlElement("w:tcBorders")
            # schema order: tcBorders comes right after tcW / gridSpan / vMerge
            anchor = [c for c in tc_pr if c.tag in (
                qn("w:tcW"), qn("w:gridSpan"), qn("w:vMerge"))]
            if anchor:
                anchor[-1].addnext(borders)
            else:
                tc_pr.insert(0, borders)
        rule = OxmlElement(f"w:{side}")
        for k, v in (("val", "single"), ("sz", str(size)), ("space", "0"), ("color", "000000")):
            rule.set(qn(f"w:{k}"), v)
        borders.append(rule)


def write_cell(cell, text: str, italic=False, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER):
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.0
    for i, line in enumerate(text.split("\n")):
        if i:
            p.add_run().add_break()
        run = p.add_run(line)
        run.italic, run.bold, run.font.size = italic, bold, Pt(10)


def add_table(doc, title: str, header: list[str], rows: list, note: str, first_col_cm: float = 5.6,
              bold_rows: set[int] = frozenset()):
    """rows: lists of cell texts; a plain string is a panel heading spanning the whole row.
    bold_rows: positions in rows to print in bold."""
    p = doc.add_paragraph(title)
    p.paragraph_format.keep_with_next = True

    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    tbl_w = t._tbl.tblPr.find(qn("w:tblW"))
    tbl_w.set(qn("w:type"), "dxa")
    # twentieths of a point
    tbl_w.set(qn("w:w"), str(int(TEXT_WIDTH_CM / 2.54 * 1440)))
    other_cm = (TEXT_WIDTH_CM - first_col_cm) / (len(header) - 1)
    for j, col in enumerate(t.columns):
        col.width = Cm(first_col_cm if j == 0 else other_cm)
        for cell in col.cells:
            cell.width = col.width

    left, center = WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER
    for j, h in enumerate(header):
        write_cell(t.rows[0].cells[j], h, align=left if j == 0 else center)
    for i, row in enumerate(rows, start=1):
        cells = t.rows[i].cells
        if isinstance(row, str):
            write_cell(cells[0].merge(cells[-1]), row, italic=True, align=left)
            continue
        for j, text in enumerate(row):
            write_cell(cells[j], text, bold=i - 1 in bold_rows,
                       align=left if j == 0 else center)
    add_rule(t.rows[0].cells, "top", 12)
    add_rule(t.rows[0].cells, "bottom", 6)
    add_rule(t.rows[-1].cells, "bottom", 12)

    n = doc.add_paragraph()
    n.paragraph_format.space_before = Pt(3)
    n.paragraph_format.space_after = Pt(18)
    n.paragraph_format.line_spacing = 1.0
    n.add_run("Notes: ").italic = True
    n.add_run(note)
    for run in n.runs:
        run.font.size = Pt(9)


# The tables

def table1(doc):
    e = read("table1_eras")
    rows = [["Games"] + [num(v, 0) for v in e["games"]],
            ["Home team wins (%)"] + [num(v) for v in e["home_win_pct"]],
            ["Three-point attempts (% of FGA)"] + [num(v)
                                                   for v in e["3pa_share_pct"]],
            ["SD of second-half margin (points)"] + [num(v, 2) for v in e["sd_second_half_margin"]]]
    add_table(doc, "Table 1: Sample and scoring environment by era", [""] + [era_header(l) for l in e.index], rows,
              "Regular-season games from ESPN, excluding neutral-site games and games tied at halftime. "
              "Three-point attempts are both teams' three-point attempts as a share of all field-goal "
              "attempts (games with box-score data, about 97%). The SD of the second-half margin is computed "
              "for close games (halftime margin within ±10 points) and measures how noisy the end of a close game is.")


def table2(doc):
    o = read("table2_down1_up1")
    rows = [["Home team wins when down 1 (%)"] + [num(v) for v in o["win_pct_down_1"]],
            ["Home team wins when up 1 (%)"] + [num(v)
                                                for v in o["win_pct_up_1"]],
            ["Difference, down 1 − up 1 (pp)"] + [num(v)
                                                  for v in o["difference_pp"]],
            ["Games down 1"] + [num(v, 0) for v in o["n_down_1"]],
            ["Games up 1"] + [num(v, 0) for v in o["n_up_1"]],
            ["Share of ±1 games at −1"] + [num(v, 2) for v in o["share_at_minus_1"]]]
    add_table(doc, "Table 2: Home win rate when down by one vs. up by one point at halftime",
              [""] + [era_header(l) for l in o.index], rows,
              "Raw shares, no regression. Without a motivational effect, being up by one should win clearly "
              "more often than being down by one. The last row checks that games are split evenly between "
              "−1 and +1 (no sorting around the cutoff).")


def table3(doc):
    m, d = read("table3a_rd_by_era"), read("table3b_era_differences")
    b, s = zip(*(est(r) for _, r in m.iterrows()))
    post, three, three_post = (est(d.loc[k]) for k in (
        "post minus B&P", "3pt minus B&P", "3pt minus post"))
    rows = ["Panel A. Jump in P(win) at a halftime margin of 0",
            ["Behind at halftime (pp)", *b], ["", *s],
            ["95% confidence interval"] +
            [f"[{num(r.ci_lo)}, {num(r.ci_hi)}]" for r in m.itertuples()],
            ["Minimum detectable effect (pp)"] + [num(v) for v in m["mde"]],
            ["Observations"] + [num(v, 0) for v in m["n"]],
            "Panel B. Difference in the jump between eras",
            ["Difference to B&P era", DASH, post[0], three[0]], [
                "", "", post[1], three[1]],
            ["Difference to post-B&P era", DASH, DASH, three_post[0]], ["", "", "", three_post[1]]]
    add_table(doc, "Table 3: Effect of trailing at halftime on winning, by era",
              [""] + [era_header(l) for l in m.index], rows,
              "Regression discontinuity at a home halftime margin of 0, one observation per game. OLS on games "
              "with a halftime margin within ±10 points, with separate linear slopes on each side of the cutoff; "
              "heteroskedasticity-robust standard errors (HC1) in parentheses. The outcome is 100 if the home team "
              "wins, so estimates are in percentage points; a positive estimate means that being just behind "
              "raises the probability of winning. The minimum detectable effect is for 80% power at the 5% level "
              f"(2.8 × SE). Panel B: difference between two era estimates from Panel A, tested with a {Z_TEST}. "
              + STARS)


def table4(doc):
    m = read("table4_mechanism")
    b, s = zip(*(est({"coef": r.points_jump, "se": r.points_se,
               "p": r.points_p}, 2) for r in m.itertuples()))
    vb, vs = zip(*(est({"coef": r.point_value, "se": r.point_value_se, "p": r.point_value_p}, 2)
                   for r in m.itertuples()))
    rows = ["Panel A. Effort: jump in the second-half points margin",
            ["Behind at halftime (points)", *b], ["", *s],
            ["Observations"] + [num(v, 0) for v in m["n"]],
            "Panel B. Noise: value of one halftime point",
            ["One point further ahead at halftime (pp)", *vb], ["", *vs],
            ["Relative to B&P era (%)"] + [num(v, 0) for v in m["point_value_vs_bp_pct"]]]
    add_table(doc, "Table 4: Mechanism: effort and noise", [""] + [era_header(l) for l in m.index], rows,
              "Panel A: same regression discontinuity as in Table 3 (linear, ±10 points, robust standard errors in "
              "parentheses), with the home team's second-half points margin (final margin minus halftime margin) as "
              "the outcome; a positive estimate means that teams just behind outscore their opponent after halftime. "
              "Panel B: increase in the home team's probability of winning per additional point of halftime margin, "
              "i.e. the average of the two slopes in the regression of Table 3 (robust standard errors in "
              "parentheses). Noisier second halves make each point worth less; the last row gives the value of a "
              "point relative to the B&P era. " + STARS)


def table5(doc):
    s = read("table5_short_line", index=[0, 1])
    cols = ["regular-line seasons", "short-line seasons", "short minus regular"]
    rows = []
    for outcome, label, d in (("win", "Jump in P(win) (pp)", 1),
                              ("second_half_margin", "Jump in second-half margin (points)", 2)):
        b, se = zip(*(est(s.loc[(outcome, c)], d) for c in cols))
        rows += [[label, *b], ["", *se]]
    add_table(doc, "Table 5: Short three-point line within the B&P era",
              ["", "Regular-line\nseasons",
                  "Short-line seasons\n1994/95–1996/97", "Difference"], rows,
              "B&P-era seasons 1993/94–2008/09 only. The NBA shortened the three-point line to 22 feet from 1994/95 "
              "to 1996/97, which temporarily raised three-point attempts from 11.7% to about 20% of field-goal "
              "attempts. Same regression discontinuity as in Table 3 (linear, ±10 points), estimated separately "
              f"for regular-line ({num(s.loc[('win', cols[0]), 'n'], 0)} games) and short-line seasons "
              f"({num(s.loc[('win', cols[1]), 'n'], 0)} games); robust standard errors in parentheses. The difference "
              f"is tested with a {Z_TEST}. " + STARS)


def table6(doc):
    r = read("table6_robustness", index=[0, 1])
    eras = list(dict.fromkeys(r.index.get_level_values(0)))
    rows = ["Panel A. Window and functional form"]
    for check, label in (("linear ±5", "Linear, ±5 points"), ("linear ±10 (main)", "Linear, ±10 points (main)"),
                         ("linear ±15", "Linear, ±15 points"), ("cubic ±10",
                                                                "Cubic, ±10 points (B&P-style)"),
                         ("rdrobust", "rdrobust, local linear")):
        b, s = zip(*(est(r.loc[(e, check)]) for e in eras))
        rows += [[label, *b], ["", *s]]
    rows.append("Panel B. Timing within the three-point era")
    for check, label in (("without COVID seasons", "Without COVID seasons"),
                         ("2014/15-2018/19 only", "2014/15–2018/19 only"),
                         ("2019/20-2024/25 only", "2019/20–2024/25 only")):
        b, s = est(r.loc[(eras[2], check)])
        rows += [[label, DASH, DASH, b], ["", "", "", s]]
    bandwidths = ", ".join(num(r.loc[(e, "rdrobust"), "bw"]) for e in eras)
    add_table(doc, "Table 6: Robustness of the jump in P(win)", [""] + [era_header(l) for l in eras], rows,
              "Each cell is the jump in the probability of winning (percentage points) at a halftime margin of 0, "
              "standard errors in parentheses. Rows in Panel A vary the window around the cutoff and the polynomial; "
              "the cubic is shown only for comparability with B&P, since high-order polynomials "
              "are not recommended in RD designs (Gelman & Imbens, 2019). Panel B uses the main specification (linear, ±10 points); the COVID "
              "seasons are 2019/20 (bubble) and 2020/21 (empty arenas). " + STARS)


def tableA1(doc):
    s = read("tableA1_seasons")
    header = ["Season", "Games", "Home wins\n(%)", "3PA\n(% of FGA)", "SD second-\nhalf margin",
              "Jump in P(win)\n(pp)"]
    rows = [[f"{t - 1}/{str(t)[-2:]}", num(r["games"], 0), num(r["home_win_pct"]), num(r["3pa_share_pct"]),
             num(r["sd_second_half_margin"], 2), " ".join(est(r))] for t, r in s.iterrows()]
    add_table(doc, "Table A1: Season-level data", header, rows,
              "Columns as in Table 1. The last column is the main regression discontinuity (Table 3) estimated "
              "season by season, robust standard error in parentheses; single seasons are imprecise (SE ≈ 7 pp). "
              + STARS + " Short-line seasons (1994/95–1996/97) in bold.", first_col_cm=1.8,
              bold_rows={i for i, t in enumerate(s.index) if SHORT_LINE[0] <= t <= SHORT_LINE[1]})


def write_tables(out_dir: Path = OUT_DIR):
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name, style.font.size = FONT, Pt(10)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = Cm(2.5)
    for build in (table1, table2, table3, table4, table5, table6, tableA1):
        build(doc)
    doc.save(out_dir / "tables.docx")


if __name__ == "__main__":
    write_tables()
    print(f"Tables saved to {OUT_DIR / 'tables.docx'}")
